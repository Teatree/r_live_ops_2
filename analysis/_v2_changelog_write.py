"""Workbook writer for the v2 config changelog. Imported by _build_v2_changelog.py."""
import os, collections
import openpyxl
from openpyxl.utils import get_column_letter as CL
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ARIAL = lambda **k: Font(name='Arial', **k)
F_BAR = 'FF1F2430'; F_HDR = 'FFD9D9D9'
F_ADD = 'FFE2EFDA'; F_CUT = 'FFFCE4E4'; F_CFG = 'FFFFF2CC'; F_NEW = 'FFCFE2F3'
THIN  = Side(style='thin', color='FFD0D0D0')
BOX   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

KIND_FILL = {'Envelope added': F_ADD, 'ToF ticket added': F_ADD, 'Reward increased': F_ADD,
             'Reward reduced': F_CUT, 'Envelope removed': F_CUT,
             'New column / helper': F_NEW, 'Working calculation': F_NEW,
             'Config parameter': F_CFG}


def _setup(ws, widths):
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[CL(i)].width = w


def _bar(ws, r, text, ncol):
    for c in range(1, ncol + 1):
        ws.cell(r, c).fill = PatternFill('solid', fgColor=F_BAR)
    ws.cell(r, 1, text).font = ARIAL(size=11, bold=True, color='FFFFFFFF')


def _hdr(ws, r, labels):
    for i, t in enumerate(labels, start=1):
        c = ws.cell(r, i, t)
        c.font = ARIAL(size=10, bold=True)
        c.fill = PatternFill('solid', fgColor=F_HDR)
        c.border = BOX
        c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)


def write(rows, per_source, why, flags, out):
    wb = openpyxl.Workbook()

    # ------------------------------------------------------------------ 1. SUMMARY
    ws = wb.active
    ws.title = 'Summary'
    _setup(ws, [26, 30, 16, 11, 10, 9, 9, 8, 82])
    ws.cell(1, 1, 'LiveOps v2 config changes - what was edited, and why').font = ARIAL(size=14, bold=True)
    ws.cell(2, 1, 'COLLECTIONS_UNDER_NEW_CALENDAR (latest_1).xlsx   |   %d authored changes across '
                  '%d source configs   |   compared inside each BASE sheet used range; formula '
                  'cells skipped; blank treated as 0' % (len(rows), len(per_source))
            ).font = ARIAL(size=9, italic=True, color='FF666666')
    ws.cell(3, 1, 'Every "Why" column below is INFERRED from the diff - a reading of what a group '
                  'of edits does together. None of it was told to me. See the Flags tab first.'
            ).font = ARIAL(size=9, italic=True, color='FFB06000')
    r = 5
    _bar(ws, r, 'BY SOURCE', 9); r += 1
    _hdr(ws, r, ['Source', 'What it is', 'Sheets', 'Envelopes', 'Tickets', 'Cuts', 'Raises',
                 'Other', 'Why this was probably done  (inferred)']); r += 1
    for src in sorted(per_source, key=lambda k: -per_source[k]['n']):
        d = per_source[src]
        mine = [x for x in rows if x['source'] == src]
        cnt = collections.Counter(x['kind'] for x in mine)
        other = len(mine) - cnt['Envelope added'] - cnt['ToF ticket added'] \
                          - cnt['Reward reduced'] - cnt['Reward increased']
        vals = [src, d['what'], d['sheet'], cnt['Envelope added'], cnt['ToF ticket added'],
                cnt['Reward reduced'], cnt['Reward increased'], other, why.get(src, '')]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(r, i, v)
            c.font = ARIAL(size=10, bold=(i == 1))
            c.border = BOX
            c.alignment = Alignment(vertical='top', wrap_text=(i in (2, 9)))
            if i == 4 and v: c.fill = PatternFill('solid', fgColor=F_ADD)
            if i == 6 and v: c.fill = PatternFill('solid', fgColor=F_CUT)
        ws.row_dimensions[r].height = 76
        r += 1
    tot = collections.Counter(x['kind'] for x in rows)
    ws.cell(r, 1, 'TOTAL').font = ARIAL(size=10, bold=True)
    for i, k in [(4, 'Envelope added'), (5, 'ToF ticket added'),
                 (6, 'Reward reduced'), (7, 'Reward increased')]:
        ws.cell(r, i, tot[k]).font = ARIAL(size=10, bold=True)
    ws.cell(r, 8, len(rows) - sum(tot[k] for k in ('Envelope added', 'ToF ticket added',
                                                   'Reward reduced', 'Reward increased'))
            ).font = ARIAL(size=10, bold=True)
    r += 2
    _bar(ws, r, 'THE PATTERN, IN ONE LINE', 9); r += 1
    ws.cell(r, 1, '%d envelopes and %d ToF tickets were ADDED across the calendar, and %d rewards '
                  'were CUT to pay for them against only %d raised. This is a swap, not a giveaway: '
                  'the collection feature is being funded out of the existing booster and coin '
                  'budget, ladder by ladder, mostly on the same rung.'
            % (tot['Envelope added'], tot['ToF ticket added'], tot['Reward reduced'],
               tot['Reward increased'])).font = ARIAL(size=10)
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=9)
    ws.cell(r, 1).alignment = Alignment(vertical='top', wrap_text=True)

    # ------------------------------------------------------------------ 2. FLAGS
    ws2 = wb.create_sheet('Flags')
    _setup(ws2, [52, 26, 112])
    ws2.cell(1, 1, 'Things worth checking before this ships').font = ARIAL(size=14, bold=True)
    ws2.cell(2, 1, 'Found while diffing. The first one is the expensive kind: it fails silently.'
             ).font = ARIAL(size=9, italic=True, color='FF666666')
    r = 4
    _bar(ws2, r, 'FLAGS', 3); r += 1
    _hdr(ws2, r, ['What', 'Source', 'Detail']); r += 1
    for what, src, detail in flags:
        for i, v in enumerate([what, src, detail], start=1):
            c = ws2.cell(r, i, v)
            c.font = ARIAL(size=10, bold=(i == 1))
            c.border = BOX
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if i == 1:
                c.fill = PatternFill('solid', fgColor=F_CUT)
        ws2.row_dimensions[r].height = 66
        r += 1

    # ------------------------------------------------------------------ 3. ALL CHANGES
    ws3 = wb.create_sheet('All changes')
    _setup(ws3, [26, 30, 22, 24, 12, 12, 20, 9, 60])
    ws3.cell(1, 1, 'Every authored change, one row each').font = ARIAL(size=14, bold=True)
    ws3.cell(2, 1, 'Sorted by source, then by cell. "Base" is the live config, "v2" the redesign.'
             ).font = ARIAL(size=9, italic=True, color='FF666666')
    r = 4
    _bar(ws3, r, 'CHANGES', 9); r += 1
    _hdr(ws3, r, ['Source', 'Block', 'Rung / row', 'Field', 'Base', 'v2', 'Change', 'Cell', 'Note'])
    r += 1
    for x in sorted(rows, key=lambda z: (z['source'], int(z['cell'].split(':')[0][1:] if
                    z['cell'][1:].split(':')[0].isdigit() else 0), z['cell'])):
        vals = [x['source'], x['block'], x['rung'], x['field'], x['base'], x['v2'],
                x['kind'], x['cell'], x['note']]
        for i, v in enumerate(vals, start=1):
            c = ws3.cell(r, i, v)
            c.font = ARIAL(size=10)
            c.border = BOX
            c.alignment = Alignment(vertical='top', wrap_text=(i == 9))
            if i == 7:
                c.fill = PatternFill('solid', fgColor=KIND_FILL.get(x['kind'], 'FFFFFFFF'))
        r += 1
    ws3.freeze_panes = 'A6'
    ws3.auto_filter.ref = 'A5:I%d' % (r - 1)

    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    return out
