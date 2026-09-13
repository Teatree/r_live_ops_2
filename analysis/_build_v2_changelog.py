#!/usr/bin/env python3
"""
_build_v2_changelog.py - every authored change on every _v2 config sheet, vs its base.

WHAT IT COMPARES. Each `<X>` / `<X>_v2` pair, cell by cell, but ONLY inside the BASE sheet's used
range. That is the same rule V2Diff.gs paints red with, and it matters: a _v2 sheet carries engine
spill blocks and helper columns to the right of the base's range, and comparing those would report
the simulator's own output as an authored change.

TWO EXCLUSIONS, both deliberate:

  * FORMULA CELLS. A formula's cached value moves with the data behind it, so a diff on values
    would report a recalculation as an edit. Skipped on either side.

  * BLANK vs 0. The config-sheet house style says "0 (not blank) for empty numeric cells", and the
    _v2 sheets follow it while several base sheets do not. Treating those as changes buried the
    real edits: Rainbow Maker alone reported 616 differences of which 584 were a blank becoming a
    zero. Blank and 0 are the same statement about a reward grid, so they compare equal here.
    Total across the workbook: 1,516 raw differences -> 349 real ones.

SOURCE NAMES, NOT SHEET CODES (user, 2026-09-14). `Ki_v2` is not a thing anyone discusses; Kite
Festival is. The map below is the same one CAL_LABEL / the config-sheet docs use.

Output: reports/LiveOps_v2_config_changelog.xlsx
"""
import os, sys, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC  = os.path.join(ROOT, 'workbooks', 'COLLECTIONS_UNDER_NEW_CALENDAR (latest_1).xlsx')
OUT  = os.path.join(ROOT, 'reports', 'LiveOps_v2_config_changelog.xlsx')

# sheet code -> (real source name, what it is)
SOURCE = {
    'c_saga':  ('Core Saga progression', 'Level-completion reward ladder (always-on)'),
    'c_day':   ('Daily Gift',            'Login streak ladder (always-on)'),
    'SP':      ('Season Pass',           '30-tier season track, free + paid'),
    'SP_lb':   ('Season Pass Leaderboard','Dream Pass challenge pot'),
    'RR':      ('River Rush',            'Removed from cal_new'),
    'HH':      ('Hatchling Hideaway',    'Collection event, 4 days'),
    'J':       ('Jigsaw Puzzle',         'Collection event, 3 days'),
    'BB':      ("Bomb's Ballet",         'Collection event, 3 days'),
    'Ph':      ('Photoshoot',            'Token-shop collection event, 3 days'),
    'Ki':      ('Kite Festival',         'Opt-in league leaderboard'),
    'NS':      ('Night Sky',             'Daily streak event'),
    'F':       ('Flock Flurry',          'Team event'),
    'Race':    ('Level Race',            'Leaderboard event'),
    'TaD':     ('Target Day',            'Leaderboard + milestone event'),
    'RM_1st':  ('Rainbow Maker',         'Milestone event, 1st-half config'),
    'RM_2nd':  ('Rainbow Maker',         'Milestone event, 2nd-half config'),
}
# (base, v2, variant label shown beside the source name)
PAIRS = [
    ('c_saga','c_saga_v2',''), ('c_day','c_day_v2',''), ('SP','SP_v2',''), ('SP_lb','SP_lb_v2',''),
    ('RR','RR_v2',''), ('HH','HH_v2',''), ('J','J_v2',''), ('BB','BB_v2',''), ('Ph','Ph_v2',''),
    ('Ki','Ki_v2',''), ('NS','NS_v2','weekend ladder'), ('NS','NS_weekday_v2','weekday ladder'),
    ('F','F_v2',''), ('Race','Race_v2',''), ('TaD','TaD_v2',''),
    ('RM_1st','RM_1st_v2','1st half'), ('RM_2nd','RM_2nd_v2','2nd half'),
]
PACK_COLS = ['1-star Dly','2-star Dly','3-star Dly','4-star Dly','5-star Dly','6-star Dly']
CURRENCY  = ['Coins','HC Reward','SPT','SPT x2','Red','Chuck','Bomb','Slingshot','Shuffle','Comet',
             'Unlimited Lives','Unlimited Red','Unlimited Chuck','Unlimited Bomb','COOP Token','Avatar']

ARIAL  = lambda **k: Font(name='Arial', **k)
F_BAR  = 'FF1F2430'; F_HDR = 'FFD9D9D9'
F_ADD  = 'FFE2EFDA'; F_CUT = 'FFFCE4E4'; F_CFG = 'FFFFF2CC'; F_NEW = 'FFCFE2F3'
THIN   = Side(style='thin', color='FFD0D0D0')
BOX    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def norm(x):
    """Blank and 0 are the same statement; a numeric string is the number it spells."""
    if x is None:
        return 0
    if isinstance(x, bool):
        return x
    if isinstance(x, str):
        s = x.strip()
        if s == '':
            return 0
        try:
            f = float(s)
            return int(f) if f == int(f) else f
        except ValueError:
            return s
    if isinstance(x, float) and x == int(x):
        return int(x)
    return x


def is_formula(v):
    return (isinstance(v, str) and v.startswith('=')) or hasattr(v, 'text')


def header_for(ws, r, c):
    """Nearest string above in the same column - the block's header row."""
    for rr in range(r - 1, 0, -1):
        v = ws.cell(rr, c).value
        if isinstance(v, str) and v.strip():
            return v.strip().replace('\n', ' ')
    return 'col%d' % c


def block_for(ws, r):
    """Nearest BAR above: a row with text in column A and nothing beside it."""
    for rr in range(r, 0, -1):
        a = ws.cell(rr, 1).value
        if isinstance(a, str) and a.strip():
            rest = [ws.cell(rr, cc).value for cc in range(2, 8)]
            if not any(x not in (None, '') for x in rest):
                return a.strip()
    return ''


def rung_for(ws, r):
    """What the changed row IS - the rung / tier / item it names."""
    a = ws.cell(r, 1).value
    if isinstance(a, (int, float)):
        return str(int(a) if float(a) == int(a) else a)
    if isinstance(a, str) and a.strip():
        return a.strip()
    for rr in range(r, 0, -1):
        v = ws.cell(rr, 1).value
        if isinstance(v, str) and v.strip():
            return '%s +%d' % (v.strip(), r - rr)
    return 'row %d' % r


def classify(field, base, v2):
    """(change type, note) for one cell."""
    num_b = isinstance(base, (int, float)) and not isinstance(base, bool)
    num_v = isinstance(v2, (int, float)) and not isinstance(v2, bool)
    if field in PACK_COLS and num_b and num_v and base == 0 and v2 > 0:
        return 'Envelope added', 'Card-collection supply: this rung now pays a %s envelope.' % field.replace(' Dly', '')
    if field == 'ToF_Ticket' and num_b and num_v and base == 0 and v2 > 0:
        return 'ToF ticket added', 'Mighty Doors entry currency seeded on this rung.'
    if field in PACK_COLS and num_b and num_v and base > 0 and v2 == 0:
        return 'Envelope removed', 'Envelope taken back off this rung.'
    if field in CURRENCY and num_b and num_v:
        if v2 < base:
            return 'Reward reduced', 'Cut from %s to %s - pays for what was added elsewhere on the ladder.' % (base, v2)
        if v2 > base:
            return 'Reward increased', 'Raised from %s to %s.' % (base, v2)
    if not num_v and v2 != 0:
        return 'New column / helper', 'New authored column or working block on the _v2 sheet.'
    # A cell under no header at all is scratch: a working calculation typed beside the config
    # rather than part of it. Worth listing (it shows what was being worked out) but it is not a
    # config change, and mislabelling it as one inflates the count.
    if re.match(r'^col\d+$', field or ''):
        return 'Working calculation', 'Scratch working typed beside the config, not a config value.'
    return 'Config parameter', 'Changed from %s to %s.' % (base, v2)


def main():
    if not os.path.exists(SRC):
        sys.exit('missing workbook: %s' % SRC)
    wv = openpyxl.load_workbook(SRC, data_only=True)
    wf = openpyxl.load_workbook(SRC, data_only=False)

    rows, per_source = [], {}
    for base, v2, variant in PAIRS:
        if base not in wv.sheetnames or v2 not in wv.sheetnames:
            continue
        B, V, BF, VF = wv[base], wv[v2], wf[base], wf[v2]
        name, what = SOURCE.get(base, (base, ''))
        label = name + ((' (%s)' % variant) if variant else '')
        nr = min(B.max_row, V.max_row)
        nc = min(B.max_column, V.max_column)
        found = 0
        for r in range(1, nr + 1):
            for c in range(1, nc + 1):
                if is_formula(BF.cell(r, c).value) or is_formula(VF.cell(r, c).value):
                    continue
                a, b = norm(B.cell(r, c).value), norm(V.cell(r, c).value)
                if a == b:
                    continue
                field = header_for(V, r, c)
                kind, note = classify(field, a, b)
                rows.append({'source': label, 'what': what, 'sheet': '%s -> %s' % (base, v2),
                             'block': block_for(V, r), 'rung': rung_for(V, r), 'field': field,
                             'base': a, 'v2': b, 'kind': kind, 'note': note,
                             'cell': '%s%d' % (openpyxl.utils.get_column_letter(c), r)})
                found += 1
        per_source.setdefault(label, {'what': what, 'sheet': '%s -> %s' % (base, v2), 'n': 0})
        per_source[label]['n'] += found
    # COLLAPSE SCRATCH RUNS. Photoshoot's token ramp is 44 cells across three rows; listing each
    # one buries the 300 real edits under an artefact. One row per (source, rung, kind) run, with
    # the cell count, says the same thing and stays readable.
    out, i = [], 0
    while i < len(rows):
        r = rows[i]
        if r['kind'] != 'Working calculation':
            out.append(r); i += 1; continue
        j = i
        while (j + 1 < len(rows) and rows[j+1]['kind'] == 'Working calculation'
               and rows[j+1]['source'] == r['source'] and rows[j+1]['rung'] == r['rung']):
            j += 1
        n = j - i + 1
        if n > 1:
            r = dict(r)
            r['field'] = '%s .. %s (%d cells)' % (rows[i]['field'], rows[j]['field'], n)
            r['base'] = ''; r['v2'] = '%s .. %s' % (rows[i]['v2'], rows[j]['v2'])
            r['cell'] = '%s:%s' % (rows[i]['cell'], rows[j]['cell'])
        out.append(r); i = j + 1
    # recount per source off the collapsed list
    per_source = {}
    for r in out:
        d = per_source.setdefault(r['source'], {'what': r['what'], 'sheet': r['sheet'], 'n': 0})
        d['n'] += 1
    return out, per_source



if __name__ == '__main__':
    import collections
    from _v2_changelog_notes import WHY, FLAGS
    from _v2_changelog_write import write
    rows, per_source = main()
    print('%d real changes across %d source configs' % (len(rows), len(per_source)))
    for k, v in collections.Counter(r['kind'] for r in rows).most_common():
        print('   %-22s %d' % (k, v))
    missing = [s for s in per_source if s not in WHY]
    if missing:
        print('NO "why" NOTE FOR: ' + ', '.join(missing))
    print('written ' + write(rows, per_source, WHY, FLAGS, OUT))
