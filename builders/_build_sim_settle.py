# Builds sim_settle_v1.xlsx — the OPTIONAL authored panel behind the daily view's day placement
# (D60, 2026-10-14; EcoGainsSim_Daily.gs settleConfig_). Import it as a sheet named 'sim_settle'.
#
#   source        a CATEGORY_ORDER name, exactly
#   settle_days   days after the grid cell on which that source's rank payout lands. The engine's
#                 default is 1 for every DAILY_LASTDAY family (a 1-day calendar cell is a 2-date
#                 event, the board settles on the next date) and 0 for everything else; a blank
#                 cell KEEPS the default, a number overrides it for that source only.
#   lag_0..lag_2  how a settled payout is spread over +0 / +1 / +2 further days, normalised by
#                 the engine. All blank = identity (the whole payout on settle day). This is the
#                 slot for change #4 of the approved plan: fit it from the grant-vs-balance lag
#                 data when it arrives; until then leave it blank - the totals cannot identify it.
#
# The sheet is read by HEADER (trimmed, case-insensitive), so columns can be reordered. The daily
# view runs the measured defaults whether or not this sheet exists; it only has to exist to
# override them. No value here is read by the 33-day sim, the 7-day view (earn day) or PBP.
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DISPLAY = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'display')

# the rank families, in CATEGORY_ORDER order — the engine's DAILY_LASTDAY set
RANK_FAMILIES = ['Bomb Challenge', 'Chuck Challenge', 'Flock Flurry', 'Kite Festival', 'Level Race',
                 'Red Challenge', 'Target Day', 'Team Event', 'Team Race', 'Flash Race']
HDR = ['source', 'settle_days', 'lag_0', 'lag_1', 'lag_2', 'note']

fill = lambda rgb: PatternFill('solid', fgColor=rgb)
thin = Side(style='thin')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
F_INPUT, F_HDR = 'FFFFF2CC', 'FFD9D9D9'

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'sim_settle'
ws.sheet_view.showGridLines = False
for col, w in zip('ABCDEF', (22, 12, 8, 8, 8, 70)):
    ws.column_dimensions[col].width = w

for j, h in enumerate(HDR, 1):
    c = ws.cell(1, j, h)
    c.font = Font(name='Arial', size=10, bold=True)
    c.fill = fill(F_HDR)
    c.border = BORDER
    c.alignment = Alignment(horizontal='center')

for i, src in enumerate(RANK_FAMILIES, 2):
    ws.cell(i, 1, src).font = Font(name='Arial', size=10)
    # settle_days 1 = the measured default written out, so the sheet documents what runs
    for j, v in ((2, 1), (3, 1), (4, 0), (5, 0)):
        c = ws.cell(i, j, v)
        c.font = Font(name='Arial', size=10)
        c.fill = fill(F_INPUT)
        c.border = BORDER
        c.alignment = Alignment(horizontal='center')
    ws.cell(i, 6, 'rank payout lands the date after the grid cell (default); lag 1/0/0 = whole payout on that date').font = \
        Font(name='Arial', size=9, color='FF808080')

notes = [
    'HOW TO READ THIS SHEET (D60, 2026-10-14)',
    'settle_days: days after a source\'s calendar cell on which its rank payout reaches the balance. The engine default is 1 for every rank family (a 1-day cell is a 2-date event; measured: Team Race Sunday closes pay on Monday in all ten cells) and 0 for every other source. Blank = keep the default.',
    'lag_0 / lag_1 / lag_2: optional spread of a settled payout over +0 / +1 / +2 days (normalised). Leave at 1/0/0 until the grant-vs-balance lag data exists to fit it — the daily totals alone cannot identify this width (a symmetric smear fits them as well as a forward one).',
    'A source not listed here runs its default. Add a row for any CATEGORY_ORDER name to override it; a collection or always-on source can be given a settle too, but nothing measured supports that today.',
    'Applies to the daily view only (balance day). The 7-day view sums the EARN day and ignores this sheet; the 33-day sim and PBP never read it.',
    'Measured effect of the default (harness/_probe_daily_alloc_toggles.js, wb15): allocation error 19.6% -> 17.8% alone; with the Season Pass lane and River Rush on-grid fixes, 12.1%.',
]
for i, t in enumerate(notes):
    c = ws.cell(len(RANK_FAMILIES) + 4 + i, 1, t)
    c.font = Font(name='Arial', size=9, bold=(i == 0), color='FF000000' if i == 0 else 'FF808080')

out = os.path.join(DISPLAY, 'sim_settle_v1.xlsx')
wb.save(out)
print('written sim_settle_v1.xlsx -', len(RANK_FAMILIES), 'rank families at settle_days 1, lag 1/0/0')
