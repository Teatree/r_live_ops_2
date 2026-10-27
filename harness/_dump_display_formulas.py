# -*- coding: utf-8 -*-
"""Dump the DISPLAY xlsx files the builders emit - values AND formulas - into one JSON the SimFill
harness reads (D65, 2026-10-27).

    python harness/_dump_display_formulas.py        -> harness/_display_formulas.json

Why a separate dump: _dump_mockdata.py captures the live workbook's VALUES, and a workbook that
has been filled carries no ECOGAINS_ formula anywhere - so no mockdata dump can say where the
builders put the anchors. The builders can. Each display file is loaded with its formulas intact;
the harness mounts the values as a sheet, runs simFillLayoutAnchors_ over them and asserts that
the anchors it derives from the LABELS are exactly the cells and formulas the builder wrote.
Re-run after changing a builder; the gate then tells you whether the layout rule still reads it.

Formula cells are blanked in `values` (openpyxl has no cached value for a formula it wrote) and
listed under `formulas` as {"r,c": "=..."} (1-based), the same keying the harness mock uses.
"""
import io
import json
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DISPLAY = os.path.join(ROOT, 'display')
OUT = os.path.join(ROOT, 'harness', '_display_formulas.json')

# file -> the sheet name the engine knows it by (the xlsx may carry the builder's own title)
FILES = [
    ('EcoGainsSim_HC_v4.xlsx', None),
    ('EcoGainsSim_Daily_v3.xlsx', None),
    ('ToF_v1.xlsx', None),
    ('EcoGainsSim_PlybyPly_v6.xlsx', None),
]


def cell_value(v):
    if v is None:
        return ''
    if isinstance(v, (int, float, bool)):
        return v
    return str(v)


def main():
    out = {}
    for fname, alias in FILES:
        path = os.path.join(DISPLAY, fname)
        if not os.path.exists(path):
            print('MISSING display file:', fname)
            continue
        wb = openpyxl.load_workbook(path, data_only=False)
        for ws in wb.worksheets:
            name = alias or ws.title
            values, formulas = [], {}
            max_r, max_c = ws.max_row, ws.max_column
            for r in range(1, max_r + 1):
                row = []
                for c in range(1, max_c + 1):
                    v = ws.cell(r, c).value
                    is_formula = (isinstance(v, str) and v.startswith('=')) or \
                                 (v is not None and type(v).__name__ == 'ArrayFormula')
                    if is_formula:
                        text = v if isinstance(v, str) else getattr(v, 'text', str(v))
                        formulas['%d,%d' % (r, c)] = text
                        row.append('')
                    else:
                        row.append(cell_value(v))
                values.append(row)
            out[name] = {'file': fname, 'values': values, 'formulas': formulas,
                         'max_rows': max_r, 'max_cols': max_c}
            n_eco = sum(1 for f in formulas.values() if 'ECOGAINS' in f.upper())
            print('%-28s <- %-30s %4d x %-3d  formulas %4d  ECOGAINS %d' % (name, fname, max_r, max_c, len(formulas), n_eco))
    with io.open(OUT, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False)
    print('written', OUT)


if __name__ == '__main__':
    main()
